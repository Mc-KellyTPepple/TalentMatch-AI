"""
TalentMatch AI - FastAPI application (v2).

Endpoints
---------
GET  /                 web interface
GET  /health           instant liveness probe (Render health check)
GET  /status           engine / feed / AI configuration (no secrets)
POST /analyze          CV upload -> profile, career roles, live jobs, prep
POST /api/analyze      alias of /analyze

Design notes
------------
* Nothing heavy is loaded at import time, so the port opens immediately and
  Render's health check passes while the service boots.
* CV analysis is CPU-bound, so it runs in a worker thread; /health always
  answers instantly, even during an analysis.
* CV text is processed in memory only and is never stored or logged.
"""

from __future__ import annotations

import io
import os
import re
import threading
import time
import traceback
from collections import defaultdict, deque
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Deque, Dict

from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.concurrency import run_in_threadpool

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
TEMPLATE_DIR = BASE_DIR / "templates"

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_CV_CHARS = 30000
MIN_CV_CHARS = 40

_START = time.perf_counter()


def log(message: str) -> None:
    print(f"[TalentMatch {time.perf_counter() - _START:8.2f}s] {message}", flush=True)


# ---------------------------------------------------------------------------
# Lifespan
# ---------------------------------------------------------------------------

@asynccontextmanager
async def lifespan(_: FastAPI):
    log("TalentMatch AI started.")
    # Warm the live-vacancy cache in the background; never blocks start-up.
    try:
        from live_jobs import store
        store.refresh_async()
    except Exception as exc:  # noqa: BLE001
        log(f"Live job warm-up skipped: {exc!r}")
    yield
    log("TalentMatch AI shutting down.")


app = FastAPI(
    title="TalentMatch AI",
    description="Match a CV to career roles and current live vacancies.",
    version="2.0.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["GET", "POST", "OPTIONS"],
    allow_headers=["*"],
)

if STATIC_DIR.exists():
    app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")

templates = Jinja2Templates(directory=str(TEMPLATE_DIR)) if TEMPLATE_DIR.exists() else None


# ---------------------------------------------------------------------------
# Text extraction
# ---------------------------------------------------------------------------

def extract_text_from_upload(filename: str, data: bytes) -> str:
    name = str(filename or "").lower()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")

    if name.endswith(".pdf"):
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                try:
                    reader.decrypt("")
                except Exception:  # noqa: BLE001
                    raise ValueError("This PDF is password protected. Please upload an unlocked copy.")
            pages = []
            for page in reader.pages[:15]:
                try:
                    text = page.extract_text() or ""
                except Exception:  # noqa: BLE001
                    continue
                if text.strip():
                    pages.append(text)
            return "\n".join(pages)
        except ValueError:
            raise
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"Could not read the PDF file ({type(exc).__name__}).")

    if name.endswith(".docx"):
        try:
            from docx import Document

            document = Document(io.BytesIO(data))
            parts = [p.text for p in document.paragraphs if (p.text or "").strip()]
            # Many CV templates keep content (skills, experience) inside tables.
            for table in document.tables:
                for row in table.rows:
                    seen = set()
                    for cell in row.cells:
                        cell_text = (cell.text or "").strip()
                        if cell_text and cell_text not in seen:
                            seen.add(cell_text)
                            parts.append(cell_text)
            return "\n".join(parts)
        except Exception as exc:  # noqa: BLE001
            raise ValueError(f"Could not read the DOCX file ({type(exc).__name__}).")

    raise ValueError("Unsupported CV format. Please upload a PDF, DOCX or TXT file.")


def clean_resume_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in str(text or "").splitlines()]
    return "\n".join(line for line in lines if line)[:MAX_CV_CHARS].strip()


# ---------------------------------------------------------------------------
# Tiny in-memory rate limiter (protects the Groq key and the free instance)
# ---------------------------------------------------------------------------

_RATE_WINDOW = 600  # seconds
_RATE_LIMIT = int(os.getenv("ANALYSES_PER_10_MIN", "15"))
_hits: Dict[str, Deque[float]] = defaultdict(deque)
_hits_lock = threading.Lock()


def _client_ip(request: Request) -> str:
    forwarded = request.headers.get("x-forwarded-for", "")
    return (forwarded.split(",")[0].strip() or (request.client.host if request.client else "unknown"))


def _check_rate(request: Request) -> None:
    ip = _client_ip(request)
    now = time.time()
    with _hits_lock:
        window = _hits[ip]
        while window and now - window[0] > _RATE_WINDOW:
            window.popleft()
        if len(window) >= _RATE_LIMIT:
            raise HTTPException(status_code=429, detail="Too many analyses from this connection. Please wait a few minutes and try again.")
        window.append(now)
        if len(_hits) > 5000:  # bound memory
            for key in [k for k, v in _hits.items() if not v]:
                _hits.pop(key, None)


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------

@app.get("/", response_class=HTMLResponse)
async def home(request: Request):
    if templates is not None:
        try:
            return templates.TemplateResponse(request, "index.html", {})
        except Exception as exc:  # noqa: BLE001
            log(f"Template rendering failed: {exc!r}")
    return HTMLResponse("<h1>TalentMatch AI</h1><p>API is running. POST a CV to /analyze.</p>")


@app.get("/health")
async def health():
    return {"status": "healthy", "service": "TalentMatch AI", "version": app.version}


@app.get("/status")
async def status():
    from groq_reranker import groq_status
    from live_jobs import store
    from roles_kb import ROLES

    return {
        "service": "TalentMatch AI",
        "version": app.version,
        "roles_in_library": len(ROLES),
        "ai_reranking": groq_status(),
        "live_jobs": store.status(),
    }


async def _read_cv(file: UploadFile) -> str:
    if not file or not file.filename:
        raise HTTPException(status_code=400, detail="No CV file was supplied.")
    data = await file.read(MAX_UPLOAD_BYTES + 1)
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded CV is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="The CV is larger than the 10 MB limit.")
    try:
        text = clean_resume_text(await run_in_threadpool(extract_text_from_upload, file.filename, data))
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    if len(text) < MIN_CV_CHARS:
        raise HTTPException(
            status_code=400,
            detail="No readable text was found in this CV. If it is a scanned image or photo, "
                   "please upload a text-based PDF, a DOCX or a TXT version.",
        )
    return text


def _truthy(value: str) -> bool:
    return str(value).strip().lower() in {"1", "true", "yes", "on"}


@app.post("/analyze")
async def analyze_cv(
    request: Request,
    file: UploadFile = File(...),
    top_k: int = Form(10),
    country: str = Form(""),
    remote_only: str = Form("false"),
    only_eligible: str = Form("false"),
    use_ai: str = Form("true"),
):
    _check_rate(request)
    started = time.perf_counter()
    resume_text = await _read_cv(file)

    from ranking_engine import analyze_cv as run_analysis

    try:
        result = await run_in_threadpool(
            run_analysis,
            resume_text,
            max(1, min(int(top_k or 10), 10)),
            country,
            _truthy(remote_only),
            _truthy(only_eligible),
            _truthy(use_ai),
        )
    except Exception as exc:  # noqa: BLE001
        log(f"CV analysis failed: {exc!r}")
        print(traceback.format_exc(), flush=True)
        raise HTTPException(status_code=500, detail="TalentMatch AI could not complete the CV analysis. Please try again.")

    return {
        "success": True,
        "filename": file.filename,
        "resume_characters": len(resume_text),
        "processing_time_seconds": round(time.perf_counter() - started, 3),
        **result,
    }


@app.post("/api/analyze")
async def api_analyze_cv(
    request: Request,
    file: UploadFile = File(...),
    top_k: int = Form(10),
    country: str = Form(""),
    remote_only: str = Form("false"),
    only_eligible: str = Form("false"),
    use_ai: str = Form("true"),
):
    return await analyze_cv(request, file, top_k, country, remote_only, only_eligible, use_ai)


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("app:app", host="0.0.0.0", port=int(os.environ.get("PORT", "8000")), workers=1)
